from fastapi import FastAPI, HTTPException, Query
from fastapi.responses import HTMLResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field
from typing import List, Optional
from pathlib import Path
from datetime import datetime
from io import BytesIO
import re, sqlite3
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR=Path(__file__).resolve().parent.parent; DATA_DIR=BASE_DIR/'data'; DATA_DIR.mkdir(exist_ok=True); DB_PATH=DATA_DIR/'urenregistratie.db'; LEGACY_WORKBOOK=DATA_DIR/'urenregistratie.xlsx'
app=FastAPI(title='UrenApp'); app.mount('/static',StaticFiles(directory=str(BASE_DIR/'app'/'static')),name='static')
DAYS=['Maandag','Dinsdag','Woensdag','Donderdag','Vrijdag','Zaterdag','Zondag']
class WorkLine(BaseModel):
    day:str; job:str=Field(min_length=1,max_length=120); hours:float=Field(ge=0,le=24); free_hours:float=Field(default=0,ge=0,le=24); note:Optional[str]=Field(default='',max_length=250)
class TemporaryWorkerLine(BaseModel):
    day:str; name:str=Field(min_length=2,max_length=60); hours:float=Field(gt=0,le=24); job:Optional[str]=Field(default='',max_length=120); note:Optional[str]=Field(default='',max_length=250)
class AbsenceLine(BaseModel):
    day:str; category:str=Field(min_length=2,max_length=40); hours:float=Field(gt=0,le=24); note:Optional[str]=Field(default='',max_length=250)
class ProjectCreate(BaseModel):
    name:str=Field(min_length=2,max_length=120)
class Submission(BaseModel):
    name:str=Field(min_length=2,max_length=60); week:int=Field(ge=1,le=53); zero_hours_contract:bool=False; lines:List[WorkLine]; temporary_workers:List[TemporaryWorkerLine]=Field(default_factory=list); absences:List[AbsenceLine]=Field(default_factory=list)
def period_for_week(w): return ((w-1)//4)+1
def safe_sheet_name(n): return (re.sub(r'[\\/*?:\[\]]','-',n).strip() or 'Medewerker')[:31]
def db():
    c=sqlite3.connect(DB_PATH,timeout=30); c.row_factory=sqlite3.Row; c.execute('PRAGMA foreign_keys=ON'); return c
def init_db():
    c=db(); c.executescript('''PRAGMA journal_mode=WAL; CREATE TABLE IF NOT EXISTS submissions(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL,name_key TEXT NOT NULL,week INTEGER NOT NULL,period INTEGER NOT NULL,zero_hours_contract INTEGER NOT NULL DEFAULT 0,is_temporary INTEGER NOT NULL DEFAULT 0,submitted_at TEXT NOT NULL,UNIQUE(name_key,week)); CREATE TABLE IF NOT EXISTS work_lines(id INTEGER PRIMARY KEY AUTOINCREMENT,submission_id INTEGER NOT NULL,day TEXT NOT NULL,job TEXT NOT NULL,work_hours REAL NOT NULL DEFAULT 0,free_hours REAL NOT NULL DEFAULT 0,note TEXT NOT NULL DEFAULT '',FOREIGN KEY(submission_id) REFERENCES submissions(id) ON DELETE CASCADE); CREATE TABLE IF NOT EXISTS temporary_work_lines(id INTEGER PRIMARY KEY AUTOINCREMENT,submission_id INTEGER NOT NULL,day TEXT NOT NULL,name TEXT NOT NULL,work_hours REAL NOT NULL DEFAULT 0,job TEXT NOT NULL DEFAULT '',note TEXT NOT NULL DEFAULT '',FOREIGN KEY(submission_id) REFERENCES submissions(id) ON DELETE CASCADE); CREATE TABLE IF NOT EXISTS projects(id INTEGER PRIMARY KEY AUTOINCREMENT,name TEXT NOT NULL UNIQUE COLLATE NOCASE,created_at TEXT NOT NULL); CREATE TABLE IF NOT EXISTS absence_lines(id INTEGER PRIMARY KEY AUTOINCREMENT,submission_id INTEGER NOT NULL,day TEXT NOT NULL,category TEXT NOT NULL,hours REAL NOT NULL DEFAULT 0,note TEXT NOT NULL DEFAULT '',FOREIGN KEY(submission_id) REFERENCES submissions(id) ON DELETE CASCADE);''');
    if not any(r[1]=='is_temporary' for r in c.execute('PRAGMA table_info(submissions)').fetchall()): c.execute('ALTER TABLE submissions ADD COLUMN is_temporary INTEGER NOT NULL DEFAULT 0')
    c.commit(); c.close(); migrate_legacy_excel_if_needed()
def migrate_legacy_excel_if_needed():
    if not LEGACY_WORKBOOK.exists(): return
    c=db()
    if c.execute('SELECT COUNT(*) FROM submissions').fetchone()[0]>0: c.close(); return
    try:
        wb=load_workbook(LEGACY_WORKBOOK,data_only=True)
        for sn in wb.sheetnames:
            if sn in {'Overzicht','Week-overzicht','Periode-overzicht'}: continue
            ws=wb[sn]; grouped={}
            for r in range(2,ws.max_row+1):
                week,day,job=ws.cell(r,2).value,ws.cell(r,3).value,ws.cell(r,4).value
                if not week or not day or not job: continue
                grouped.setdefault(int(week),[]).append((str(day),str(job),float(ws.cell(r,5).value or 0),float(ws.cell(r,6).value or 0),str(ws.cell(r,7).value or ''),str(ws.cell(r,8).value or datetime.now().strftime('%d-%m-%Y %H:%M'))))
            for week,lines in grouped.items():
                cur=c.execute('INSERT OR IGNORE INTO submissions(name,name_key,week,period,zero_hours_contract,submitted_at) VALUES(?,?,?,?,?,?)',(sn,sn.strip().lower(),week,period_for_week(week),1 if any(x[3]>0 for x in lines) else 0,lines[-1][5]))
                if cur.lastrowid:
                    for x in lines:c.execute('INSERT INTO work_lines(submission_id,day,job,work_hours,free_hours,note) VALUES(?,?,?,?,?,?)',(cur.lastrowid,x[0],x[1],x[2],x[3],x[4]))
        c.commit()
    finally:c.close()
def save_submission(s):
    name=s.name.strip(); key=name.lower(); valid=[x for x in s.lines if x.hours!=0 or x.free_hours!=0]; temporary=[x for x in s.temporary_workers if x.hours>0]; absences=[x for x in s.absences if x.hours>0]
    if not valid and not temporary and not absences: raise HTTPException(400,'Vul minimaal 1 urenregel, uitzendkracht of afwezigheid in.')
    c=db()
    try:
        c.execute('BEGIN IMMEDIATE'); old=c.execute('SELECT id FROM submissions WHERE name_key=? AND week=?',(key,s.week)).fetchone()
        if old:c.execute('DELETE FROM submissions WHERE id=?',(old['id'],))
        stamp=datetime.now().strftime('%d-%m-%Y %H:%M'); cur=c.execute('INSERT INTO submissions(name,name_key,week,period,zero_hours_contract,is_temporary,submitted_at) VALUES(?,?,?,?,?,?,?)',(name,key,s.week,period_for_week(s.week),1 if s.zero_hours_contract else 0,0,stamp)); wt=ft=tt=at=0.0
        for x in valid:
            w=float(x.hours); f=float(x.free_hours) if s.zero_hours_contract else 0.0; wt+=w; ft+=f; c.execute('INSERT INTO work_lines(submission_id,day,job,work_hours,free_hours,note) VALUES(?,?,?,?,?,?)',(cur.lastrowid,x.day,x.job.strip(),w,f,(x.note or '').strip()))
        for x in temporary:
            h=float(x.hours); tt+=h; c.execute('INSERT INTO temporary_work_lines(submission_id,day,name,work_hours,job,note) VALUES(?,?,?,?,?,?)',(cur.lastrowid,x.day,x.name.strip(),h,(x.job or '').strip(),(x.note or '').strip()))
        for x in absences:
            h=float(x.hours); at+=h; c.execute('INSERT INTO absence_lines(submission_id,day,category,hours,note) VALUES(?,?,?,?,?)',(cur.lastrowid,x.day,x.category.strip(),h,(x.note or '').strip()))
        c.commit(); return {'period':period_for_week(s.week),'work_total':round(wt,2),'free_total':round(ft,2),'temporary_total':round(tt,2),'absence_total':round(at,2)}
    except Exception:c.rollback(); raise
    finally:c.close()
def style_header(ws):
    fill=PatternFill('solid',fgColor='1F4E78')
    for x in ws[1]:x.font=Font(color='FFFFFF',bold=True);x.fill=fill;x.alignment=Alignment(vertical='center')
    ws.row_dimensions[1].height=22
def autosize(ws):
    d={}
    for row in ws.iter_rows():
        for x in row:
            if x.value is not None:d[x.column]=max(d.get(x.column,0),len(str(x.value)))
    for col,w in d.items():ws.column_dimensions[get_column_letter(col)].width=min(max(w+2,10),36)
def fetch_summary_rows():
    c=db(); r=c.execute('''SELECT s.id,s.submitted_at,s.name,s.name_key,s.week,s.period,s.zero_hours_contract,COALESCE((SELECT SUM(w.work_hours) FROM work_lines w WHERE w.submission_id=s.id),0) work_hours,COALESCE((SELECT SUM(w.free_hours) FROM work_lines w WHERE w.submission_id=s.id),0) free_hours,COALESCE((SELECT GROUP_CONCAT(DISTINCT w.job) FROM work_lines w WHERE w.submission_id=s.id),'') work_addresses,COALESCE((SELECT SUM(t.work_hours) FROM temporary_work_lines t WHERE t.submission_id=s.id),0) temporary_hours,COALESCE((SELECT GROUP_CONCAT(DISTINCT t.name) FROM temporary_work_lines t WHERE t.submission_id=s.id),'') temporary_names,COALESCE((SELECT SUM(a.hours) FROM absence_lines a WHERE a.submission_id=s.id),0) absence_hours,COALESCE((SELECT GROUP_CONCAT(DISTINCT a.category) FROM absence_lines a WHERE a.submission_id=s.id),'') absence_types FROM submissions s ORDER BY s.week DESC,s.name COLLATE NOCASE''').fetchall();c.close();return r
def build_excel():
    wb=Workbook();ov=wb.active;ov.title='Overzicht';ov.append(['Ingediend op','Naam','Periode','Week','Projecturen','Uitzenduren','Afwezigheidsuren','Vrije uren','Totaal uren','Uitzendkrachten','Soort afwezigheid','0-urencontract']);style_header(ov);ov.freeze_panes='A2'
    wo=wb.create_sheet('Week-overzicht');wo.append(['Week','Periode','Naam','Projecturen','Uitzenduren','Afwezigheidsuren','Vrije uren','Totaal uren','Ingediend op']);style_header(wo);wo.freeze_panes='A2'
    po=wb.create_sheet('Periode-overzicht');po.append(['Periode','Week','Naam','Projecturen','Uitzenduren','Afwezigheidsuren','Vrije uren','Totaal uren']);style_header(po);po.freeze_panes='A2'
    summary=fetch_summary_rows()
    for r in summary:
        w=round(float(r['work_hours']),2);u=round(float(r['temporary_hours']),2);a=round(float(r['absence_hours']),2);f=round(float(r['free_hours']),2);t=round(w+u+a+f,2);ov.append([r['submitted_at'],r['name'],r['period'],r['week'],w,u,a,f,t,r['temporary_names'],r['absence_types'],'Ja' if r['zero_hours_contract'] else 'Nee']);wo.append([r['week'],r['period'],r['name'],w,u,a,f,t,r['submitted_at']]);po.append([r['period'],r['week'],r['name'],w,u,a,f,t])
    for ws in (ov,wo,po):autosize(ws)
    c=db()
    for emp in c.execute('SELECT DISTINCT name,name_key FROM submissions ORDER BY name COLLATE NOCASE').fetchall():
        ws=wb.create_sheet(safe_sheet_name(emp['name']));ws.append(['Periode','Week','Dag','Klus / werkadres','Werkuren','Vrije uren','Opmerking','Ingediend op']);style_header(ws);ws.freeze_panes='A2'
        rr=c.execute('''SELECT s.period,s.week,w.day,w.job,w.work_hours,w.free_hours,w.note,s.submitted_at FROM submissions s JOIN work_lines w ON w.submission_id=s.id WHERE s.name_key=? ORDER BY s.week,CASE w.day WHEN 'Maandag' THEN 1 WHEN 'Dinsdag' THEN 2 WHEN 'Woensdag' THEN 3 WHEN 'Donderdag' THEN 4 WHEN 'Vrijdag' THEN 5 WHEN 'Zaterdag' THEN 6 WHEN 'Zondag' THEN 7 ELSE 8 END,w.id''',(emp['name_key'],)).fetchall()
        for r in rr:ws.append([r['period'],r['week'],r['day'],r['job'],r['work_hours'],r['free_hours'],r['note'],r['submitted_at']])
        autosize(ws)
    tw=wb.create_sheet('Uitzendkrachten');tw.append(['Periode','Week','Dag','Naam uitzendkracht','Uren','Klus / werkadres','Opmerking','Ingevoerd door','Ingediend op']);style_header(tw);tw.freeze_panes='A2'
    for r in c.execute('''SELECT s.period,s.week,t.day,t.name,t.work_hours,t.job,t.note,s.name entered_by,s.submitted_at FROM temporary_work_lines t JOIN submissions s ON s.id=t.submission_id ORDER BY s.week,t.day,t.name COLLATE NOCASE''').fetchall():tw.append([r['period'],r['week'],r['day'],r['name'],r['work_hours'],r['job'],r['note'],r['entered_by'],r['submitted_at']])
    autosize(tw)
    aw=wb.create_sheet('Verlof en afspraken');aw.append(['Periode','Week','Dag','Medewerker','Soort','Uren','Opmerking','Ingediend op']);style_header(aw);aw.freeze_panes='A2'
    for r in c.execute('''SELECT s.period,s.week,a.day,s.name,a.category,a.hours,a.note,s.submitted_at FROM absence_lines a JOIN submissions s ON s.id=a.submission_id ORDER BY s.week,a.day,s.name COLLATE NOCASE''').fetchall():aw.append([r['period'],r['week'],r['day'],r['name'],r['category'],r['hours'],r['note'],r['submitted_at']])
    autosize(aw);c.close();out=BytesIO();wb.save(out);out.seek(0);return out
@app.on_event('startup')
def startup():init_db()
@app.get('/',response_class=HTMLResponse)
def home():return (BASE_DIR/'app'/'static'/'index.html').read_text(encoding='utf-8')
@app.get('/admin',response_class=HTMLResponse)
def admin():return (BASE_DIR/'app'/'static'/'admin.html').read_text(encoding='utf-8')
@app.post('/api/submit')
def submit_hours(s:Submission):
    if any(x.day not in DAYS for x in s.lines) or any(x.day not in DAYS for x in s.temporary_workers) or any(x.day not in DAYS for x in s.absences):raise HTTPException(400,'Ongeldige dag gevonden.')
    allowed={'Vakantie','Tandarts','Ziekenhuisafspraak','Doktersafspraak','Bijzonder verlof','Overig'}
    if any(x.category not in allowed for x in s.absences):raise HTTPException(400,'Ongeldige soort afwezigheid.')
    return {'ok':True,'message':'Week succesvol ingediend.',**save_submission(s)}
@app.get('/api/status')
def status():
    out=[]
    for r in fetch_summary_rows():
        w=round(float(r['work_hours']),2);u=round(float(r['temporary_hours']),2);a=round(float(r['absence_hours']),2);f=round(float(r['free_hours']),2);out.append({'id':r['id'],'submitted_at':r['submitted_at'],'name':r['name'],'name_key':r['name_key'],'work_addresses':r['work_addresses'],'period':r['period'],'week':r['week'],'work_hours':w,'temporary_hours':u,'temporary_names':r['temporary_names'],'absence_hours':a,'absence_types':r['absence_types'],'free_hours':f,'total_hours':round(w+u+a+f,2),'zero_hours_contract':'Ja' if r['zero_hours_contract'] else 'Nee','worker_type':'Medewerker'})
    return out
@app.delete('/api/submission/{sid}')
def delete_submission(sid:int):
    c=db();r=c.execute('SELECT name,week FROM submissions WHERE id=?',(sid,)).fetchone()
    if not r:c.close();raise HTTPException(404,'Inzending niet gevonden.')
    c.execute('DELETE FROM submissions WHERE id=?',(sid,));c.commit();c.close();return {'ok':True,'message':f"Week {r['week']} van {r['name']} is verwijderd."}
@app.delete('/api/employee')
def delete_employee(name_key:str=Query(...)):
    c=db();r=c.execute('SELECT name FROM submissions WHERE name_key=? LIMIT 1',(name_key,)).fetchone()
    if not r:c.close();raise HTTPException(404,'Medewerker niet gevonden.')
    name=r['name'];c.execute('DELETE FROM submissions WHERE name_key=?',(name_key,));c.commit();c.close();return {'ok':True,'message':f'{name} en alle weken zijn verwijderd.'}
@app.get('/api/projects')
def list_projects():
    c=db();rows=c.execute('SELECT id,name FROM projects ORDER BY name COLLATE NOCASE').fetchall();c.close();return [dict(r) for r in rows]
@app.post('/api/projects')
def add_project(project:ProjectCreate):
    name=project.name.strip();c=db()
    try:
        cur=c.execute('INSERT INTO projects(name,created_at) VALUES(?,?)',(name,datetime.now().strftime('%d-%m-%Y %H:%M')));c.commit();return {'id':cur.lastrowid,'name':name}
    except sqlite3.IntegrityError:raise HTTPException(409,'Dit project bestaat al.')
    finally:c.close()
@app.delete('/api/projects/{project_id}')
def delete_project(project_id:int):
    c=db();r=c.execute('SELECT name FROM projects WHERE id=?',(project_id,)).fetchone()
    if not r:c.close();raise HTTPException(404,'Project niet gevonden.')
    c.execute('DELETE FROM projects WHERE id=?',(project_id,));c.commit();c.close();return {'ok':True,'message':f"Project {r['name']} is verwijderd."}
@app.get('/api/excel')
def excel_download():return StreamingResponse(build_excel(),media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="urenregistratie.xlsx"'})
@app.get('/health')
def health():return {'ok':True,'storage':'sqlite','excel':'generated-live'}
